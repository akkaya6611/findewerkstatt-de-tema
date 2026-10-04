import openpyxl
import re
import json
import sys

def parse_workshops():
    xlsx_path = r'C:\Users\Serkan\.gemini\antigravity\brain\a66a60b5-6687-4a17-b9ce-240284dfd061\.user_uploaded\media_1791060353110.xlsx'
    wb = openpyxl.load_workbook(xlsx_path)
    sheet = wb.active

    workshops = []
    skipped = 0

    known_brands = {
        'volkswagen': ['volkswagen', 'vw'],
        'bmw': ['bmw'],
        'mercedes-benz': ['mercedes', 'mercedes-benz', 'daimler'],
        'audi': ['audi'],
        'opel': ['opel'],
        'ford': ['ford'],
        'porsche': ['porsche'],
        'skoda': ['skoda', 'škoda'],
        'seat': ['seat', 'cupra'],
        'renault': ['renault', 'dacia'],
        'peugeot': ['peugeot'],
        'toyota': ['toyota', 'lexus'],
        'hyundai': ['hyundai'],
        'kia': ['kia'],
        'fiat': ['fiat', 'abarth', 'alfa romeo'],
        'volvo': ['volvo'],
        'mazda': ['mazda'],
        'nissan': ['nissan'],
    }

    for r in range(2, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        if not any(vals):
            continue

        search_cat = str(vals[0] or '').strip()
        city = str(vals[1] or '').strip()
        district = str(vals[2] or '').strip()
        plz_raw = vals[3]
        name = str(vals[4] or '').strip()
        category = str(vals[5] or '').strip()
        rating_raw = str(vals[6] or '').strip()
        reviews_raw = str(vals[7] or '').strip()
        phone_raw = str(vals[8] or '').strip()
        addr_raw = str(vals[9] or '').strip()
        website_raw = str(vals[10] or '').strip()
        image_raw = str(vals[11] or '').strip()
        maps_raw = str(vals[12] or '').strip()

        # Skip invalid / empty rows
        if name in ('', 'N/A') or addr_raw in ('', 'N/A') or len(name) < 2:
            skipped += 1
            continue

        # PLZ formatting (ensure 5 digits string)
        plz = ''
        if plz_raw and str(plz_raw) != 'N/A':
            plz = str(plz_raw).strip()
            if plz.isdigit():
                plz = plz.zfill(5)

        # Rating parsing (e.g. '4,9' -> 4.9)
        rating = None
        if rating_raw and rating_raw != 'N/A':
            r_str = rating_raw.replace(',', '.')
            try:
                rating = round(float(r_str), 1)
            except ValueError:
                rating = None

        # Review count parsing (e.g. '179 Berichte' -> 179, '1.500 Berichte' -> 1500)
        review_count = 0
        if reviews_raw and reviews_raw != 'N/A':
            digits_only = re.sub(r'[^0-9]', '', reviews_raw.split()[0].replace('.', ''))
            if digits_only.isdigit():
                review_count = int(digits_only)

        # Street address extraction (extract street + number without repeating PLZ / city)
        street = addr_raw
        match_plz_split = re.split(r',\s*\d{5}\b', addr_raw)
        if match_plz_split and match_plz_split[0].strip():
            street = match_plz_split[0].strip()

        # District cleaning
        clean_district = ''
        if district and district != 'N/A' and district.lower() != city.lower():
            clean_district = district

        # Phone cleaning
        phone = ''
        if phone_raw and phone_raw != 'N/A':
            phone = phone_raw

        # Website cleaning
        website = ''
        if website_raw and website_raw != 'N/A':
            website = website_raw

        # Image cleaning
        image_url = ''
        is_real_photo = False
        if image_raw and image_raw != 'N/A':
            image_url = image_raw
            if not any(sub in image_url for sub in ['w36-h36', 'ALV-U', 'ACg8oc']):
                is_real_photo = True

        # Lat / Lng coordinates extraction
        lat = None
        lng = None
        m = re.search(r'!3d([0-9.]+)!4d([0-9.]+)', maps_raw)
        if m:
            lat = float(m.group(1))
            lng = float(m.group(2))
        else:
            m2 = re.search(r'@([0-9.]+),([0-9.]+)', maps_raw)
            if m2:
                lat = float(m2.group(1))
                lng = float(m2.group(2))

        # Badges
        is_verified = (rating is not None and rating >= 4.5 and review_count >= 10)
        is_master = ('meister' in name.lower() or 'meister' in category.lower())
        emergency_24h = ('24h' in name.lower() or 'notdienst' in name.lower() or 'notdienst' in category.lower())

        # Detected brands
        matched_brands = []
        name_cat_lower = (name + ' ' + category).lower()
        for brand_slug, aliases in known_brands.items():
            for alias in aliases:
                # check whole word match
                if re.search(r'\b' + re.escape(alias) + r'\b', name_cat_lower):
                    matched_brands.append(brand_slug)
                    break

        item = {
            'name': name,
            'search_category': search_cat,
            'category': category,
            'city': city,
            'district': clean_district,
            'plz': plz,
            'street': street,
            'full_address': addr_raw,
            'phone': phone,
            'website': website,
            'rating': rating,
            'review_count': review_count,
            'lat': lat,
            'lng': lng,
            'maps_url': maps_raw,
            'image_url': image_url,
            'is_real_photo': is_real_photo,
            'is_verified': is_verified,
            'is_master': is_master,
            'emergency_24h': emergency_24h,
            'brands': matched_brands,
        }
        workshops.append(item)

    out_file = r'c:\Users\Serkan\.gemini\antigravity\scratch\findewerkstatt-de-tema\data\workshops.json'
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(workshops, f, ensure_ascii=False, indent=2)

    print(f'Successfully parsed {len(workshops)} workshops (skipped {skipped} invalid rows).')
    print(f'Saved to {out_file}')

if __name__ == '__main__':
    parse_workshops()
